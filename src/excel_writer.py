"""Writes a batch of downloaded symbol DataFrames into one .xlsx workbook.

One sheet per symbol, per docs/DESIGN.md §12. A symbol that failed to fetch
simply isn't in the input dict, so it never gets an (empty) sheet.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def sanitize_sheet_name(name: str) -> str:
    """Strip Excel-illegal characters and truncate to Excel's 31-char sheet name limit."""
    cleaned = _INVALID_SHEET_CHARS.sub("_", name)
    return cleaned[:31] or "Sheet"


def write_workbook(data: dict[str, pd.DataFrame], output_path: Path) -> None:
    """Write `data` (sheet name -> DataFrame) to a single .xlsx at output_path.

    Column widths are auto-sized and the header row is bolded for readability.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    used_names: dict[str, int] = {}
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for raw_name, df in data.items():
            sheet_name = sanitize_sheet_name(raw_name)
            if sheet_name in used_names:
                used_names[sheet_name] += 1
                sheet_name = sanitize_sheet_name(f"{sheet_name}_{used_names[sheet_name]}")
            else:
                used_names[sheet_name] = 0

            df.to_excel(writer, sheet_name=sheet_name, index=False)
            _format_sheet(writer.sheets[sheet_name], df)


def _format_sheet(sheet, df: pd.DataFrame) -> None:
    """Bold the header row and auto-size each column to its widest cell."""
    for col_idx, column in enumerate(df.columns, start=1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        max_len = max([len(str(column))] + [len(str(v)) for v in df[column].astype(str)])
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)
